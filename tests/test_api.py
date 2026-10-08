from io import BytesIO

from openpyxl import load_workbook


def test_interface_and_static_assets_are_served(client):
    page = client.get("/")
    assert page.status_code == 200
    assert "Nova vistoria técnica" in page.text
    assert "10</span> Pendências encontradas" in page.text
    assert "REGISTRO FOTOGRÁFICO" not in page.text
    assert client.get("/favicon.ico").status_code == 200
    assert client.get("/static/favicon.svg").status_code == 200
    assert client.get("/static/css/style.css").status_code == 200
    script = client.get("/static/js/app.js")
    assert script.status_code == 200
    assert "function collectPayload" in script.text


def sample_payload(status="Rascunho"):
    return {
        "unidade_id": "U0014",
        "status": status,
        "respostas": {
            "identificacao": {
                "data_vistoria": "2026-10-08",
                "horario_inicio": "08:00",
                "horario_termino": "10:15",
                "equipe_ipq": "Equipe Alfa",
                "responsavel_local": "Responsável local",
            },
            "pontos_camera": [{
                "local_ambiente": "Entrada principal",
                "tipo": "Dome",
                "tombamento": "TJ-001",
                "altura_aproximada": 3.2,
                "infra": True,
                "cabo": False,
                "energia": True,
                "observacoes": "Sem obstruções",
            }],
            "infraestrutura": {
                "padrao_pvc_preto": True,
                "padrao_galvanizada": False,
                "padrao_pvc_branco": False,
                "itens": [{"item": "Tubulação", "pvc_preto": True, "galvanizado": True, "pvc_branco": False}],
            },
            "cabeamento": [{"item": "Cabo", "resposta": "Existente", "observacao": None}],
            "equipamentos": [{
                "item": "Rack", "existente": "Sim", "quantidade_necessaria": 1,
                "condicao_observacao": "Bom estado", "fixacao_adequada": "Sim", "necessita_adequacao": "Não",
            }],
            "condicoes_rack": [{"item": "Ventilação adequada", "resposta": "N/A", "observacao": None}],
            "infraestrutura_eletrica": [{"item": "Aterramento disponível", "resposta": "Não", "observacao": "Adequar"}],
            "certificacao": [{"item": "Certificação realizada", "resposta": None, "observacao": None}],
            "pendencias": [{"descricao": "Aterramento", "local": "Rack", "impacto": "Segurança", "solucao_necessaria": "Instalar"}],
            "solucoes": [{"descricao": "Novo aterramento", "responsavel_dependencia": "TJCE", "prazo_observacao": "30 dias"}],
            "conclusao": {"situacao": None, "observacoes_finais": "Aguardando adequação"},
            "responsaveis": {"ipq": {"nome": "Ana", "cargo": "Técnica", "assinatura": None, "data": "2026-10-08"}, "tjce": {}},
        },
    }


def test_catalog_preserves_all_units_and_duplicates(client):
    response = client.get("/api/unidades")
    assert response.status_code == 200
    units = response.json()
    assert len(units) == 238
    duplicates = [unit for unit in units if unit["name"] == "ACOPIARA - RESIDÊNCIA"]
    assert [unit["display_name"] for unit in duplicates] == ["ACOPIARA - RESIDÊNCIA (1)", "ACOPIARA - RESIDÊNCIA (2)"]
    assert any(unit["name"] == "AQUIRAZ - FÓRUM" and unit["id"] == "U0014" for unit in units)


def test_create_read_update_and_list_visit(client):
    created = client.post("/api/vistorias", json=sample_payload())
    assert created.status_code == 201
    visit = created.json()
    assert visit["id"] == 1
    assert visit["unidade_nome"] == "AQUIRAZ - FÓRUM"
    assert visit["status"] == "Rascunho"

    loaded = client.get("/api/vistorias/1")
    assert loaded.status_code == 200
    assert loaded.json()["respostas"]["pontos_camera"][0]["tombamento"] == "TJ-001"

    payload = sample_payload("Concluída")
    payload["respostas"]["conclusao"]["situacao"] = "APTO COM RESSALVAS"
    payload["respostas"]["identificacao"]["equipe_ipq"] = "Equipe revisada"
    updated = client.put("/api/vistorias/1", json=payload)
    assert updated.status_code == 200
    assert updated.json()["id"] == 1
    assert updated.json()["status"] == "Concluída"

    listing = client.get("/api/vistorias", params={"unidade_id": "U0014", "data_inicio": "2026-10-01"})
    assert listing.status_code == 200
    assert listing.json()[0]["equipe_ipq"] == "Equipe revisada"


def test_completed_visit_requires_conclusion(client):
    response = client.post("/api/vistorias", json=sample_payload("Concluída"))
    assert response.status_code == 422
    assert "conclusão" in str(response.json()).lower()


def test_multiple_visits_for_same_unit_are_not_overwritten(client):
    assert client.post("/api/vistorias", json=sample_payload()).status_code == 201
    second = sample_payload()
    second["respostas"]["identificacao"]["data_vistoria"] = "2026-10-09"
    assert client.post("/api/vistorias", json=second).status_code == 201
    records = client.get("/api/vistorias", params={"unidade_id": "U0014"}).json()
    assert {item["id"] for item in records} == {1, 2}


def test_excel_individual_and_filtered_export(client):
    client.post("/api/vistorias", json=sample_payload())
    response = client.get("/api/vistorias/1/exportar")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == [
        "Vistorias", "Pontos de Câmera", "Infraestrutura", "Cabeamento", "Rack e Equipamentos",
        "Infraestrutura Elétrica", "Certificação", "Pendências", "Soluções", "Conclusão", "Responsáveis",
    ]
    assert workbook["Pontos de Câmera"]["E2"].value == "Entrada principal"
    assert workbook["Infraestrutura"]["D2"].value == "Tubulação"
    assert workbook["Pendências"]["D2"].value == 1
    assert workbook["Vistorias"].freeze_panes == "A2"
    assert workbook["Vistorias"].auto_filter.ref

    filtered = client.get("/api/exportar", params={"unidade_id": "U0014", "data_inicio": "2026-10-08", "data_fim": "2026-10-08"})
    assert filtered.status_code == 200
    filtered_book = load_workbook(BytesIO(filtered.content))
    assert filtered_book["Vistorias"].max_row == 2


def test_invalid_unit_and_not_found(client):
    payload = sample_payload()
    payload["unidade_id"] = "INVALID"
    assert client.post("/api/vistorias", json=payload).status_code == 422
    assert client.get("/api/vistorias/999").status_code == 404

