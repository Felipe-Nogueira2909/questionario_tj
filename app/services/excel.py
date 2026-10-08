from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Vistoria


SHEET_NAMES = [
    "Vistorias",
    "Pontos de Câmera",
    "Infraestrutura",
    "Cabeamento",
    "Rack e Equipamentos",
    "Infraestrutura Elétrica",
    "Certificação",
    "Pendências",
    "Soluções",
    "Conclusão",
    "Responsáveis",
]


def informed(value):
    return "Não informado" if value is None or value == "" else value


def marked(value):
    return "Marcado" if value is True else "Não marcado"


def add_rows(ws, headers: list[str], rows: Iterable[list]):
    ws.append(headers)
    for row in rows:
        ws.append(row)


def base(visit: Vistoria) -> list:
    return [visit.id, visit.unidade_id, visit.unidade_nome]


def build_workbook(visits: list[Vistoria]) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {name: wb.create_sheet(name) for name in SHEET_NAMES}

    summary_rows = []
    point_rows = []
    infra_rows = []
    cabling_rows = []
    rack_rows = []
    electric_rows = []
    cert_rows = []
    issue_rows = []
    solution_rows = []
    conclusion_rows = []
    responsible_rows = []

    for visit in visits:
        data = visit.respostas
        ident = data.get("identificacao", {})
        conclusion = data.get("conclusao", {})
        common = base(visit)
        summary_rows.append(common + [
            visit.data_vistoria,
            informed(ident.get("horario_inicio")),
            informed(ident.get("horario_termino")),
            informed(ident.get("equipe_ipq")),
            informed(ident.get("responsavel_local")),
            informed(conclusion.get("situacao")),
            visit.status,
            visit.criado_em,
            visit.atualizado_em,
        ])
        for number, point in enumerate(data.get("pontos_camera", []), 1):
            point_rows.append(common + [number, informed(point.get("local_ambiente")), informed(point.get("tipo")),
                informed(point.get("tombamento")), informed(point.get("altura_aproximada")), marked(point.get("infra")),
                marked(point.get("cabo")), marked(point.get("energia")), informed(point.get("observacoes"))])
        infra = data.get("infraestrutura", {})
        for item in infra.get("itens", []):
            infra_rows.append(common + [item.get("item"), marked(item.get("pvc_preto")), marked(item.get("galvanizado")),
                marked(item.get("pvc_branco")), marked(infra.get("padrao_pvc_preto")),
                marked(infra.get("padrao_galvanizada")), marked(infra.get("padrao_pvc_branco"))])
        for check in data.get("cabeamento", []):
            cabling_rows.append(common + [check.get("item"), informed(check.get("resposta")), informed(check.get("observacao"))])
        for equipment in data.get("equipamentos", []):
            rack_rows.append(common + ["Equipamento", equipment.get("item"), informed(equipment.get("existente")),
                informed(equipment.get("quantidade_necessaria")), informed(equipment.get("condicao_observacao")),
                informed(equipment.get("fixacao_adequada")), informed(equipment.get("necessita_adequacao"))])
        for check in data.get("condicoes_rack", []):
            rack_rows.append(common + ["Condição do rack", check.get("item"), informed(check.get("resposta")),
                "Não informado", informed(check.get("observacao")), "Não informado", "Não informado"])
        for check in data.get("infraestrutura_eletrica", []):
            electric_rows.append(common + [check.get("item"), informed(check.get("resposta")), informed(check.get("observacao"))])
        for check in data.get("certificacao", []):
            cert_rows.append(common + [check.get("item"), informed(check.get("resposta")), informed(check.get("observacao"))])
        for number, issue in enumerate(data.get("pendencias", []), 1):
            issue_rows.append(common + [number, informed(issue.get("descricao")), informed(issue.get("local")),
                informed(issue.get("impacto")), informed(issue.get("solucao_necessaria"))])
        for number, solution in enumerate(data.get("solucoes", []), 1):
            solution_rows.append(common + [number, informed(solution.get("descricao")),
                informed(solution.get("responsavel_dependencia")), informed(solution.get("prazo_observacao"))])
        conclusion_rows.append(common + [informed(conclusion.get("situacao")), informed(conclusion.get("observacoes_finais"))])
        for role, label in (("ipq", "Responsável IPQ"), ("tjce", "Responsável TJCE / Unidade")):
            person = data.get("responsaveis", {}).get(role, {})
            responsible_rows.append(common + [label, informed(person.get("nome")), informed(person.get("cargo")),
                informed(person.get("assinatura")), informed(person.get("data"))])

    common_headers = ["ID da vistoria", "ID da unidade", "Unidade"]
    add_rows(sheets["Vistorias"], common_headers + ["Data", "Início", "Término", "Equipe IPQ", "Responsável local", "Situação", "Status", "Criação", "Última atualização"], summary_rows)
    add_rows(sheets["Pontos de Câmera"], common_headers + ["Nº", "Local / Ambiente", "Tipo", "Tombamento", "Altura aproximada (m)", "Infra", "Cabo", "Energia", "Observações"], point_rows)
    add_rows(sheets["Infraestrutura"], common_headers + ["Item", "PVC Preto", "Galvanizado", "PVC Branco", "Padrão PVC preto", "Padrão galvanizada", "Padrão PVC branco"], infra_rows)
    add_rows(sheets["Cabeamento"], common_headers + ["Verificação", "Resposta", "Observação"], cabling_rows)
    add_rows(sheets["Rack e Equipamentos"], common_headers + ["Tipo", "Item", "Existente / Resposta", "Quantidade necessária", "Condição / Observação", "Fixação adequada", "Necessita adequação"], rack_rows)
    add_rows(sheets["Infraestrutura Elétrica"], common_headers + ["Verificação", "Resposta", "Observação"], electric_rows)
    add_rows(sheets["Certificação"], common_headers + ["Verificação", "Resposta", "Resultado / Observação"], cert_rows)
    add_rows(sheets["Pendências"], common_headers + ["Nº", "Pendência / Não conformidade", "Local", "Impacto", "Solução necessária"], issue_rows)
    add_rows(sheets["Soluções"], common_headers + ["Nº", "Solução proposta", "Responsável / Dependência", "Prazo / Observação"], solution_rows)
    add_rows(sheets["Conclusão"], common_headers + ["Situação do local", "Observações finais"], conclusion_rows)
    add_rows(sheets["Responsáveis"], common_headers + ["Tipo", "Nome", "Cargo", "Assinatura", "Data"], responsible_rows)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            max_length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column)
            ws.column_dimensions[letter].width = min(max(max_length + 2, 12), 48)
            for cell in column[1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws["A"]:
            if cell.row > 1:
                cell.number_format = "0"

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream

